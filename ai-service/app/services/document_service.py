"""
Document processing service using llama-rag-sdk.

使用 llama-rag-sdk 进行文档处理：解析、分块、向量化、存储
"""
import aiofiles
import os
from pathlib import Path
from typing import Dict, Any, Optional
from datetime import datetime

from fastapi import HTTPException
from starlette import status

from app.core.logging import get_logger
from app.core.config import settings
from app.core.database import get_database
from app.models.database import DocumentModel
from app.services.sdk_adapter.config import setup_sdk_env

# 设置 SDK 环境变量
setup_sdk_env()

logger = get_logger(__name__)

# 临时文件下载目录
UPLOAD_DIR = Path(__file__).parent.parent.parent.parent / "upload_docs"
UPLOAD_DIR.mkdir(parents=True, exist_ok=True)


class DocumentProcessor:
    """
    文档处理服务（使用 llama-rag-sdk）

    SDK 处理：
    - PDF 解析（MinerU）
    - 智能分块
    - Embedding 生成
    - 存储（ChromaDB + MongoDB DocStore）
    """

    def __init__(self):
        self._rag_system = None

    @property
    def rag_system(self):
        """延迟初始化 llama-rag-sdk RAGSystem"""
        if self._rag_system is None:
            from llama_rag_sdk.rag_system import RAGSystem
            self._rag_system = RAGSystem(
                collection_name="rag_documents",
                enable_image_description=False,
                enable_summarization=True,
            )
        return self._rag_system

    async def _extract_text(self, file_path: str, file_ext: str) -> str:
        """提取文本（用于非 PDF 格式）"""
        if file_ext == ".txt":
            async with aiofiles.open(file_path, 'r', encoding='utf-8') as f:
                return await f.read()
        elif file_ext == ".md":
            async with aiofiles.open(file_path, 'r', encoding='utf-8') as f:
                return await f.read()
        elif file_ext == ".docx":
            from docx import Document
            doc = Document(file_path)
            return "\n".join([paragraph.text for paragraph in doc.paragraphs])
        elif file_ext == ".html":
            from bs4 import BeautifulSoup
            async with aiofiles.open(file_path, 'r', encoding='utf-8') as f:
                html_content = await f.read()
            soup = BeautifulSoup(html_content, 'html.parser')
            return soup.get_text()
        else:
            raise ValueError(f"Unsupported file format: {file_ext}")

    async def process_document(
        self,
        file_path: str,
        filename: str,
        kb_id: str,
        enhance: int,
        category: Optional[str] = None,
        metadata: Optional[Dict[str, Any]] = None,
        task_id: Optional[str] = None,
        chunk_config: Optional[Dict[str, Any]] = None,
        doc_id: Optional[str] = None,
        resource_id: Optional[int] = None,
    ) -> str:
        """
        处理文档：使用 llama-rag-sdk 完成解析、分块、向量化、存储

        Args:
            file_path: 文档文件路径
            filename: 原始文件名
            kb_id: 知识库 ID
            enhance: 知识增强标记
            category: 文档分类
            metadata: 额外元数据
            task_id: 任务 ID（用于进度跟踪）
            chunk_config: 自定义分块配置（SDK 暂不支持，预留）
            doc_id: 预生成的文档 ID
            resource_id: 外部系统资源 ID

        Returns:
            文档 ID
        """
        try:
            # 获取文件信息
            file_size = os.path.getsize(file_path)
            file_ext = os.path.splitext(filename)[1].lower()

            # 合并元数据
            doc_metadata = metadata.copy() if metadata else {}
            if resource_id:
                doc_metadata["resource_id"] = resource_id

            # 插入文档记录到 MongoDB
            db = await get_database()
            doc = await db.documents.find_one({"doc_id": doc_id})

            if not doc:
                doc_model = DocumentModel(
                    doc_id=doc_id,
                    filename=filename,
                    kb_id=kb_id,
                    enhance=enhance,
                    category=category,
                    size=file_size,
                    format=file_ext[1:].upper(),
                    status="processing",
                    segment_config=chunk_config,
                    metadata=doc_metadata,
                )

                result = await db.documents.insert_one(doc_model.model_dump())

                if result and result.inserted_id:
                    logger.info(
                        "insert-document",
                        doc_id=doc_id,
                        filename=filename,
                        kb_id=kb_id,
                    )
                else:
                    logger.error(f"insert-document {doc_id} failed")
                    raise HTTPException(
                        status_code=status.HTTP_500_INTERNAL_SERVER_ERROR,
                        detail="Failed to create document record"
                    )

            # 使用 llama-rag-sdk 索引文档
            # SDK 会处理：
            # 1. PDF 解析（MinerU）
            # 2. 智能分块
            # 3. Embedding 生成
            # 4. 存储（ChromaDB + MongoDB DocStore）

            if file_ext == ".pdf":
                # 传递元数据给 SDK
                await self.rag_system.index_document(
                    file_path,
                    metadata={
                        "doc_id": doc_id,
                        "kb_id": kb_id,
                        "filename": filename,
                    }
                )
            else:
                # 非 PDF 格式：先读取文本，然后添加到索引
                text_content = await self._extract_text(file_path, file_ext)

                # 手动创建文档并索引
                from llama_rag_sdk.document_parser.base import ParsedDocument, TextChunk

                parsed_doc = ParsedDocument(
                    title=filename,
                    content=text_content,
                    chunks=[
                        TextChunk(
                            text=text_content,
                            page=0,
                            chunk_index=0,
                            metadata={
                                "doc_id": doc_id,
                                "kb_id": kb_id,
                                "filename": filename,
                            }
                        )
                    ]
                )

                await self.rag_system.index_parsed_document(parsed_doc, source_path=file_path)

            # 更新文档状态
            await db.documents.update_one(
                {"doc_id": doc_id},
                {"$set": {
                    "status": "completed",
                    "processed_at": datetime.utcnow(),
                }}
            )

            logger.info(
                "Completed processing document",
                doc_id=doc_id,
                filename=filename,
                kb_id=kb_id,
            )

            return doc_id

        except Exception as e:
            logger.error(
                "Failed to process document",
                filename=filename,
                error=str(e),
                exc_info=True,
            )

            # 更新文档状态为失败
            try:
                db = await get_database()
                await db.documents.update_one(
                    {"doc_id": doc_id},
                    {"$set": {"status": "failed", "error_message": str(e)}},
                )
            except Exception as update_error:
                logger.warning(
                    "Failed to update document status to failed",
                    doc_id=doc_id,
                    error=str(update_error),
                )

            raise HTTPException(
                status_code=status.HTTP_500_INTERNAL_SERVER_ERROR,
                detail=f"Document processing failed: {str(e)}"
            )

    async def delete_document(self, doc_id: str, kb_id: str) -> None:
        """
        删除文档及其所有 chunks

        Args:
            doc_id: 文档 ID
            kb_id: 知识库 ID
        """
        try:
            db = await get_database()

            # 先查询获取该文档的所有 chunk IDs
            from app.core.chroma import chroma_db
            results = await chroma_db.query_documents(
                collection_name="doc",
                query_texts=[""],  # 空查询获取所有文档
                n_results=1000,  # 假设一个文档不超过 1000 个 chunk
                where={"doc_id": doc_id}
            )

            # 提取 chunk IDs
            chunk_ids = results.get("ids", [[]])[0] if results.get("ids") else []

            # 删除 ChromaDB 中的向量
            if chunk_ids:
                await chroma_db.delete_documents(
                    collection_name="doc",
                    ids=chunk_ids
                )

            # 删除 Elasticsearch 中的文档
            from app.core.elasticsearch import es_db
            await es_db.delete_by_query(
                index="doc",
                body={"query": {"term": {"doc_id": doc_id}}}
            )

            # 删除 MongoDB 中的 chunks
            await db.document_chunks.delete_many({"doc_id": doc_id})

            # 删除 MongoDB 中的文档记录
            await db.documents.delete_one({"doc_id": doc_id})

            logger.info("Document deleted", doc_id=doc_id, kb_id=kb_id)

        except Exception as e:
            logger.error("Failed to delete document", doc_id=doc_id, error=str(e), exc_info=True)
            raise
