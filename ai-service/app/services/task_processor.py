"""
Async document task processor with queue management.

使用 llama-rag-sdk RAGSystem 进行文档处理。
"""
import asyncio
import os
import random
import string
from datetime import datetime
from pathlib import Path
from typing import Dict, Optional, List

from app.core.logging import get_logger
from app.core.database import get_database
from app.models.database import DocumentModel, DocumentTaskModel
from app.services.dataset_faq_service import faq_processor
from app.services.thesaurus_major_service import thesaurus_major_processor
from app.services.thesaurus_sensitive_service import thesaurus_sensitive_processor
from app.services.document_service import generate_doc_id

logger = get_logger(__name__)


# ============ 文本提取工具函数 ============

async def extract_text_from_file(file_path: str, file_ext: str) -> str:
    """
    从非 PDF 文件提取文本

    Args:
        file_path: 文件路径
        file_ext: 文件扩展名（包含点号，如 .txt）

    Returns:
        提取的文本内容

    Raises:
        ValueError: 不支持的文件格式
    """
    import aiofiles
    from pathlib import Path

    if file_ext in [".txt", ".md"]:
        async with aiofiles.open(file_path, 'r', encoding='utf-8') as f:
            return await f.read()
    elif file_ext == ".docx":
        from docx import Document
        doc = Document(file_path)
        return "\n".join([paragraph.text for paragraph in doc.paragraphs])
    elif file_ext == ".html":
        from bs4 import BeautifulSoup
        async with aiofiles.open(file_path, 'r', encoding='utf-8') as f:
            html_content = await f.read()
        soup = BeautifulSoup(html_content, 'html.parser')
        return soup.get_text()
    else:
        raise ValueError(f"Unsupported file format: {file_ext}")


# ============ 任务处理器 ============


def generate_task_id() -> str:
    """Generate unique task ID."""
    timestamp = datetime.utcnow().strftime("%Y%m%d%H%M")
    random_chars = ''.join(random.choices(string.ascii_lowercase + string.digits, k=8))
    return f"document_task_{timestamp}_{random_chars}"


class DocumentTaskProcessor:
    """Background task processor for document processing.

    使用 llama-rag-sdk RAGSystem 进行文档处理，包括：
    - PDF 解析（MinerU）
    - 结构感知分块
    - Embedding 生成
    - 存储（ChromaDB + MongoDB DocStore）
    """

    def __init__(self):
        self.task_queue: asyncio.Queue = asyncio.Queue()
        self.running = False
        self.worker_task: Optional[asyncio.Task] = None
        self.active_tasks: Dict[str, bool] = {}  # task_id -> cancellation flag
        self._rag_systems: Dict[str, "RAGSystem"] = {}  # kb_id -> RAGSystem 缓存

    def _get_rag_system(self, kb_id: str):
        """
        获取或创建指定 kb_id 的 RAGSystem

        每个 kb_id 使用独立的 ChromaDB 集合（rag_documents_<kb_id>）

        Args:
            kb_id: 知识库 ID

        Returns:
            对应的 RAGSystem 实例
        """
        if kb_id not in self._rag_systems:
            from llama_rag_sdk.rag_system import RAGSystem
            logger.info(f"创建 RAGSystem for kb_id={kb_id}, 集合名=rag_documents_{kb_id}")
            self._rag_systems[kb_id] = RAGSystem(
                kb_id=kb_id,  # 自动生成集合名 rag_documents_<kb_id>
                enable_image_description=False,
                enable_summarization=True,
            )
        return self._rag_systems[kb_id]

    async def close(self):
        """关闭所有 RAGSystem 资源"""
        for rag_system in self._rag_systems.values():
            await rag_system.close()
        self._rag_systems.clear()

    async def submit_task(
        self,
        kb_id: str,
        filename: str,
        file_path: str,
        category: Optional[str] = None,
        chunk_config: Optional[Dict] = None,  # Custom chunk configuration
        doc_id: Optional[str] = None,  # Pre-generated doc_id (for immediate return)
        resource_id: Optional[int] = None  # External system resource ID
    ) -> str:
        """
        Submit a new document processing task.

        Args:
            kb_id: Knowledge base ID
            filename: Original filename
            file_path: Path to uploaded file
            category: Document category
            chunk_config: Custom chunking configuration (optional)
            doc_id: Pre-generated document ID (optional, for immediate return)
            resource_id: External system resource ID (optional)

        Returns:
            Task ID
        """
        task_id = generate_task_id()

        # Create task record in database
        db = await get_database()
        task_model = DocumentTaskModel(
            task_id=task_id,
            kb_id=kb_id,
            enhance=1,  # 保持默认值以兼容现有数据
            filename=filename,
            file_path=file_path,
            category=category,
            status="pending",
            metadata={
                "chunk_config": chunk_config,
                "resource_id": resource_id
            } if chunk_config or resource_id else {}
        )

        task_dict = task_model.model_dump()

        # Add doc_id if provided (pre-generated)
        if doc_id:
            task_dict["doc_id"] = doc_id
        if resource_id:
            task_dict["resource_id"] = resource_id

        await db.document_tasks.insert_one(task_dict)

        # Add to queue
        await self.task_queue.put({
            "task_id": task_id,
            "kb_id": kb_id,
            "filename": filename,
            "file_path": file_path,
            "category": category,
            "chunk_config": chunk_config,
            "doc_id": doc_id,  # Pass doc_id to processing
            "resource_id": resource_id
        })

        logger.info(
            f"Document task submitted: task_id={task_id}, filename={filename}, doc_id={doc_id}, resource_id={resource_id}"
        )

        return task_id

    async def re_submit_task(
        self,
        task_id: str,
        kb_id: str,
        filename: str,
        file_path: str,
        category: Optional[str] = None,
        chunk_config: Optional[Dict] = None,  # Custom chunk configuration
        doc_id: Optional[str] = None,  # Pre-generated doc_id (for immediate return)
        resource_id: Optional[int] = None  # External system resource ID):
    ):
        # Add to queue
        await self.task_queue.put({
            "task_id": task_id,
            "kb_id": kb_id,
            "filename": filename,
            "file_path": file_path,
            "category": category,
            "chunk_config": chunk_config,
            "doc_id": doc_id,  # Pass doc_id to processing
            "resource_id": resource_id
        })

        logger.info(f"re_submit_task task_id={task_id} doc_id={doc_id}")

    async def start(self):
        """Start the background task processor."""
        if self.running:
            logger.warning("Task processor already running")
            return
        
        self.running = True
        self.worker_task = asyncio.create_task(self._process_tasks())
        logger.info("Document task processor started")
    
    async def stop(self):
        """Stop the background task processor."""
        if not self.running:
            return

        self.running = False

        # Cancel all active tasks
        for task_id in list(self.active_tasks.keys()):
            await self.cancel_task(task_id)

        # Wait for worker to finish
        if self.worker_task:
            self.worker_task.cancel()
            try:
                await self.worker_task
            except asyncio.CancelledError:
                pass

        # 关闭 RAGSystem 资源
        await self.close()

        logger.info("Document task processor stopped")
    
    async def _process_tasks(self):
        """Background worker that processes tasks from the queue."""
        logger.info("Task processor worker started")
        
        while self.running:
            try:
                # Get task from queue with timeout
                try:
                    task_data = await asyncio.wait_for(
                        self.task_queue.get(),
                        timeout=1.0
                    )
                    logger.info(f"Fetched task from queue: task_id={task_data['task_id']}, task_type={task_data.get('task_type', 'document')}")
                except asyncio.TimeoutError:
                    continue
                
                task_id = task_data["task_id"]
                task_type = task_data.get("task_type", "document")
                
                # Mark task as active
                self.active_tasks[task_id] = False  # cancellation flag
                
                # Dispatch task based on type
                if task_type == "faq_vectorization":
                    await asyncio.create_task(self._execute_faq_vectorization(task_data))
                elif task_type == "thesaurus_major_vectorization":
                    await asyncio.create_task(self._execute_thesaurus_major_vectorization(task_data))
                elif task_type == "thesaurus_sensitive_vectorization":
                    await asyncio.create_task(self._execute_thesaurus_sensitive_vectorization(task_data))
                else:
                    # Default to document processing
                    await asyncio.create_task(self._execute_document_vectorization(task_data))
                
            except Exception as e:
                logger.error(f"Error in task processor loop: error={str(e)}", exc_info=True)
                await asyncio.sleep(1)

    async def get_task_status(self, task_id: str) -> Optional[Dict]:
        """
        Get task status.
        
        Args:
            task_id: Task ID
            
        Returns:
            Task status or None if not found
        """
        db = await get_database()
        task = await db.document_tasks.find_one({"task_id": task_id})
        
        if not task:
            return None
        
        task.pop("_id", None)
        
        # Format timestamps
        if task.get("created_at"):
            task["created_at"] = task["created_at"].isoformat() + "Z"
        if task.get("started_at"):
            task["started_at"] = task["started_at"].isoformat() + "Z"
        if task.get("completed_at"):
            task["completed_at"] = task["completed_at"].isoformat() + "Z"
        
        return task
    
    async def cancel_task(self, task_id: str) -> bool:
        """
        Cancel a running task.
        
        Args:
            task_id: Task ID
            
        Returns:
            True if task was cancelled, False otherwise
        """
        db = await get_database()
        
        # Get current task status
        task = await db.document_tasks.find_one({"task_id": task_id})
        
        if not task:
            return False
        
        status = task["status"]
        
        # Can only cancel pending or running tasks
        if status not in ["pending", "running"]:
            return False
        
        # If pending, update status directly
        if status == "pending":
            await db.document_tasks.update_one(
                {"task_id": task_id},
                {
                    "$set": {
                        "status": "cancelled",
                        "completed_at": datetime.utcnow()
                    }
                }
            )
            logger.info(f"Cancelled pending task: task_id={task_id}")
            return True
        
        # If running, set cancellation flag
        if task_id in self.active_tasks:
            self.active_tasks[task_id] = True  # Set cancellation flag
            logger.info(f"Cancellation requested for running task: task_id={task_id}")
            return True
        
        return False

    async def submit_faq_vectorization_task(self, faq_id: str, kb_id: str) -> str:
        """
            提交FAQ向量化任务（异步后台处理）。
            从MongoDB的faqs集合中读取该FAQ进行向量化。
        """
        task_id = f"faq_task_{datetime.utcnow().strftime('%Y%m%d%H%M%S')}_{faq_id}"

        # Create task record in database
        db = await get_database()
        insert_data = DocumentTaskModel(
            task_id=task_id,
            kb_id=kb_id,
            enhance=1,
            filename="",
            file_path="",
            doc_id=faq_id,
            category=None,
            status="pending",
            metadata={}
        ).model_dump()
        await db.document_tasks.insert_one(insert_data)

        # Add to queue for background processing
        await self.task_queue.put({
            "task_type": "faq_vectorization",
            "task_id": task_id,
            "faq_id": faq_id
        })

        logger.info(f"submit_faq_vectorization_task: task_id={task_id}")

        return task_id

    async def submit_thesaurus_major_vectorization_task(self, thesaurus_id: str, kb_id: str) -> str:
        """
            提交专业词库向量化任务（异步后台处理）。
            从MongoDB的专业词库集合中读取该专业词库的词条进行向量化。
        """
        task_id = f"major_task_{datetime.utcnow().strftime('%Y%m%d%H%M%S')}_{thesaurus_id}"

        # Create task record in database
        db = await get_database()
        insert_data = DocumentTaskModel(
            task_id=task_id,
            kb_id=kb_id,
            enhance=1,
            filename="",
            file_path="",
            doc_id=thesaurus_id,
            category=None,
            status="pending",
            metadata={}
        ).model_dump()
        await db.document_tasks.insert_one(insert_data)

        # Add to queue for background processing
        await self.task_queue.put({
            "task_type": "thesaurus_major_vectorization",
            "task_id": task_id,
            "thesaurus_id": thesaurus_id
        })

        logger.info(f"submit_thesaurus_major_vectorization_task: task_id={task_id}")

        return task_id

    async def submit_thesaurus_sensitive_vectorization_task(self, thesaurus_id: str, kb_id: str) -> str:
        """
            提交敏感词库向量化任务（异步后台处理）。
            从MongoDB的敏感词库集合中读取该敏感词库的词条进行向量化。
        """
        task_id = f"sensitive_task_{datetime.utcnow().strftime('%Y%m%d%H%M%S')}_{thesaurus_id}"

        # Create task record in database
        db = await get_database()
        insert_data = DocumentTaskModel(
            task_id=task_id,
            kb_id=kb_id,
            enhance=1,
            filename="",
            file_path="",
            doc_id=thesaurus_id,
            category=None,
            status="pending",
            metadata={}
        ).model_dump()
        await db.document_tasks.insert_one(insert_data)

        # Add to queue for background processing
        await self.task_queue.put({
            "task_type": "thesaurus_sensitive_vectorization",
            "task_id": task_id,
            "thesaurus_id": thesaurus_id
        })

        logger.info(f"submit_thesaurus_sensitive_vectorization_task: task_id={task_id}")

        return task_id

    async def _execute_document_vectorization(self, task_data: Dict):
        """Execute a single document processing task."""
        task_id = task_data["task_id"]
        db = await get_database()

        try:
            # Update task status to running
            await db.document_tasks.update_one(
                {"task_id": task_id},
                {
                    "$set": {
                        "status": "running",
                        "started_at": datetime.now()
                    }
                }
            )

            logger.info(f"Processing document task: task_id={task_id}, task_data={task_data}")

            # Process document with progress tracking
            doc_id = await self._exec_process_document(task_data)

            # Check if task was cancelled
            if self.active_tasks.get(task_id, False):
                await db.document_tasks.update_one(
                    {"task_id": task_id},
                    {
                        "$set": {
                            "status": "cancelled",
                            "completed_at": datetime.now()
                        }
                    }
                )
                logger.info(f"Document task cancelled: task_id={task_id}")
                return

            # Update task status to completed
            await db.document_tasks.update_one(
                {"task_id": task_id},
                {
                    "$set": {
                        "status": "completed",
                        "doc_id": doc_id,
                        "progress": 100.0,
                        "completed_at": datetime.now()
                    }
                }
            )

            logger.info(f"Document task completed: task_id={task_id}, doc_id={doc_id}")

        except Exception as e:
            logger.error(f"Document task failed: task_id={task_id}, error={str(e)}", exc_info=True)

            # Update task status to failed
            await db.document_tasks.update_one(
                {"task_id": task_id},
                {
                    "$set": {
                        "status": "failed",
                        "error_message": str(e),
                        "completed_at": datetime.now()
                    }
                }
            )

        finally:
            # Remove from active tasks
            self.active_tasks.pop(task_id, None)

    async def _exec_process_document(self, task_data: Dict) -> str:
        """
        使用 RAGSystem 处理文档

        流程：
        1. 创建 MongoDB 文档记录
        2. 使用 RAGSystem 索引文档
        3. 更新文档状态为完成

        Args:
            task_data: 任务数据

        Returns:
            文档 ID
        """
        task_id = task_data["task_id"]
        file_path = task_data["file_path"]
        filename = task_data["filename"]
        kb_id = task_data["kb_id"]
        category = task_data.get("category")
        doc_id = task_data.get("doc_id") or generate_doc_id(filename, kb_id)
        resource_id = task_data.get("resource_id")

        # 获取文件信息
        file_size = os.path.getsize(file_path)
        file_ext = os.path.splitext(filename)[1].lower()

        # 合并元数据
        doc_metadata = {}
        if resource_id:
            doc_metadata["resource_id"] = resource_id

        # 创建/更新 MongoDB 文档记录
        db = await get_database()
        doc = await db.documents.find_one({"doc_id": doc_id})

        if not doc:
            doc_model = DocumentModel(
                doc_id=doc_id,
                filename=filename,
                kb_id=kb_id,
                category=category,
                size=file_size,
                format=file_ext[1:].upper(),
                status="processing",
                metadata=doc_metadata,
            )

            result = await db.documents.insert_one(doc_model.model_dump())

            if result and result.inserted_id:
                logger.info(
                    "insert-document",
                    doc_id=doc_id,
                    filename=filename,
                    kb_id=kb_id,
                )
            else:
                logger.error(f"insert-document {doc_id} failed")
                raise ValueError("Failed to create document record")

        try:
            # 使用 RAGSystem 索引文档
            if file_ext == ".pdf":
                # PDF: 使用 MinerU 解析 + 结构感知分块
                await self._get_rag_system(kb_id).index_document(
                    file_path,
                    metadata={
                        "doc_id": doc_id,
                        "kb_id": kb_id,
                        "filename": filename,
                    }
                )
            else:
                # 非 PDF 格式: 先提取文本，然后创建 ParsedDocument 索引
                text_content = await extract_text_from_file(file_path, file_ext)

                from llama_rag_sdk.document_parser.base import ParsedDocument, TextChunk

                parsed_doc = ParsedDocument(
                    title=filename,
                    content=text_content,
                    chunks=[
                        TextChunk(
                            text=text_content,
                            page=0,
                            chunk_index=0,
                            metadata={
                                "doc_id": doc_id,
                                "kb_id": kb_id,
                                "filename": filename,
                            }
                        )
                    ]
                )

                await self._get_rag_system(kb_id).index_parsed_document(parsed_doc, source_path=file_path)

            # 更新文档状态为完成
            await db.documents.update_one(
                {"doc_id": doc_id},
                {"$set": {
                    "status": "completed",
                    "processed_at": datetime.utcnow(),
                }}
            )

            logger.info(
                "Document processing completed",
                doc_id=doc_id,
                filename=filename,
                kb_id=kb_id,
                task_id=task_id,
            )

            return doc_id

        except Exception as e:
            logger.error(
                "Document processing failed",
                doc_id=doc_id,
                filename=filename,
                error=str(e),
                exc_info=True,
            )

            # 更新文档状态为失败
            await db.documents.update_one(
                {"doc_id": doc_id},
                {"$set": {"status": "failed", "error_message": str(e)}},
            )

            raise

    async def _execute_faq_vectorization(self, task_data: Dict):
        """ 执行FAQ向量化任务 """
        db = await get_database()
        task_id = task_data["task_id"]

        try:
            # Update task status to running
            await db.document_tasks.update_one(
                {"task_id": task_id},
                {
                    "$set": {
                        "status": "running",
                        "started_at": datetime.now()
                    }
                }
            )

            logger.info(f"_execute_faq_vectorization task task_id={task_id}, task_data={task_data}")

            faq_id = await faq_processor.faq_vectorization(
                    task_id=task_id,
                    faq_id=task_data["faq_id"]
                )

            # Check if task was cancelled
            if self.active_tasks.get(task_id, False):
                await db.document_tasks.update_one(
                    {"task_id": task_id},
                    {
                        "$set": {
                            "status": "cancelled",
                            "completed_at": datetime.now()
                        }
                    }
                )
                logger.info(f"_execute_faq_vectorization task cancelled task_id={task_id}")
                return

            # Update task status to completed
            await db.document_tasks.update_one(
                {"task_id": task_id},
                {
                    "$set": {
                        "status": "completed",
                        "doc_id": faq_id,
                        "progress": 100.0,
                        "completed_at": datetime.now()
                    }
                }
            )

            logger.info(f"_execute_faq_vectorization task completed: task_id={task_id}, faq_id={faq_id}")

        except Exception as e:
            logger.error(f"_execute_faq_vectorization task failed: task_id={task_id}, error={str(e)}", exc_info=True)

            # Update task status to failed
            await db.document_tasks.update_one(
                {"task_id": task_id},
                {
                    "$set": {
                        "status": "failed",
                        "error_message": str(e),
                        "completed_at": datetime.now()
                    }
                }
            )

        finally:
            # Remove from active tasks
            self.active_tasks.pop(task_id, None)

    async def _execute_thesaurus_major_vectorization(self, task_data: Dict):
        """ 执行专业词库向量化任务 """
        db = await get_database()
        task_id = task_data["task_id"]

        try:
            # Update task status to running
            await db.document_tasks.update_one(
                {"task_id": task_id},
                {
                    "$set": {
                        "status": "running",
                        "started_at": datetime.now()
                    }
                }
            )

            logger.info(f"_execute_thesaurus_major_vectorization task task_id={task_id}, task_data={task_data}")

            thesaurus_id = await thesaurus_major_processor.thesaurus_major_vectorization(
                    task_id=task_id,
                    thesaurus_id=task_data["thesaurus_id"]
                )

            # Check if task was cancelled
            if self.active_tasks.get(task_id, False):
                await db.document_tasks.update_one(
                    {"task_id": task_id},
                    {
                        "$set": {
                            "status": "cancelled",
                            "completed_at": datetime.now()
                        }
                    }
                )
                logger.info(f"_execute_thesaurus_major_vectorization task cancelled task_id={task_id}")
                return

            # Update task status to completed
            await db.document_tasks.update_one(
                {"task_id": task_id},
                {
                    "$set": {
                        "status": "completed",
                        "doc_id": thesaurus_id,
                        "progress": 100.0,
                        "completed_at": datetime.now()
                    }
                }
            )

            logger.info(f"_execute_thesaurus_major_vectorization task completed: task_id={task_id}, thesaurus_id={thesaurus_id}")

        except Exception as e:
            logger.error(f"_execute_thesaurus_major_vectorization task failed: task_id={task_id}, error={str(e)}", exc_info=True)

            # Update task status to failed
            await db.document_tasks.update_one(
                {"task_id": task_id},
                {
                    "$set": {
                        "status": "failed",
                        "error_message": str(e),
                        "completed_at": datetime.now()
                    }
                }
            )

        finally:
            # Remove from active tasks
            self.active_tasks.pop(task_id, None)

    async def _execute_thesaurus_sensitive_vectorization(self, task_data: Dict):
        """ 执行敏感词库向量化任务 """
        db = await get_database()
        task_id = task_data["task_id"]

        try:
            # Update task status to running
            await db.document_tasks.update_one(
                {"task_id": task_id},
                {
                    "$set": {
                        "status": "running",
                        "started_at": datetime.now()
                    }
                }
            )

            logger.info(f"_execute_thesaurus_sensitive_vectorization task task_id={task_id}, task_data={task_data}")

            thesaurus_id = await thesaurus_sensitive_processor.thesaurus_sensitive_vectorization(
                    task_id=task_id,
                    thesaurus_id=task_data["thesaurus_id"]
                )

            # Check if task was cancelled
            if self.active_tasks.get(task_id, False):
                await db.document_tasks.update_one(
                    {"task_id": task_id},
                    {
                        "$set": {
                            "status": "cancelled",
                            "completed_at": datetime.now()
                        }
                    }
                )
                logger.info(f"_execute_thesaurus_sensitive_vectorization task cancelled task_id={task_id}")
                return

            # Update task status to completed
            await db.document_tasks.update_one(
                {"task_id": task_id},
                {
                    "$set": {
                        "status": "completed",
                        "doc_id": thesaurus_id,
                        "progress": 100.0,
                        "completed_at": datetime.now()
                    }
                }
            )

            logger.info(f"_execute_thesaurus_sensitive_vectorization task completed: task_id={task_id}, thesaurus_id={thesaurus_id}")

        except Exception as e:
            logger.error(f"_execute_thesaurus_sensitive_vectorization task failed: task_id={task_id}, error={str(e)}", exc_info=True)

            # Update task status to failed
            await db.document_tasks.update_one(
                {"task_id": task_id},
                {
                    "$set": {
                        "status": "failed",
                        "error_message": str(e),
                        "completed_at": datetime.now()
                    }
                }
            )

        finally:
            # Remove from active tasks
            self.active_tasks.pop(task_id, None)


# Global task processor instance
task_processor = DocumentTaskProcessor()
